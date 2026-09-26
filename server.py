import html
import json
import os
import re
import threading
import uuid
import webbrowser
from datetime import date, datetime, time
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HOST = "127.0.0.1"
PORT = 8000
ROOT = Path(__file__).resolve().parent
DATA_FILE = ROOT / "appointments.xlsx"
PENDING_FILE = ROOT / "appointments_pending.json"
PUBLIC_FILES = {"/", "/index.html", "/style.css"}
MAX_BODY = 10_000

LABELS = {
    "date": "Date", "time": "Time", "first_name": "First name", "last_name": "Last name",
    "gender": "Gender", "phone": "Phone number", "city": "City", "state": "State",
    "appointment_types": "Appointment type", "query": "Query",
}
REQUIRED = ["date", "time", "first_name", "last_name", "gender", "phone", "city", "state"]
GENDERS = {"male": "Male", "female": "Female", "other": "Other"}
TYPES = {"cervix": "Cervix checkup", "heart": "Heart checkup", "eye": "Eye checkup",
         "hearing": "Hearing test"}

COLUMNS = [
    ("booking_id", "Booking ID", 13),
    ("submitted_at", "Booked on", 18),
    ("date", "Appointment date", 17),
    ("time", "Time", 9),
    ("name", "Patient name", 24),
    ("phone", "Phone number", 16),
    ("gender", "Gender", 10),
    ("city", "City", 16),
    ("state", "State", 16),
    ("appointment_types", "Appointment type", 30),
    ("query", "Query", 45),
    ("contacted", "Contacted?", 12),
    ("notes", "Staff notes", 32),
]
FORMATS = {"submitted_at": "dd-mm-yyyy hh:mm", "date": "dd-mm-yyyy", "time": "hh:mm", "phone": "@"}

write_lock = threading.Lock()


def parse_form(body):
    data = parse_qs(body, keep_blank_values=True)
    record = {key: data.get(key, [""])[0].strip() for key in LABELS if key != "appointment_types"}
    record["gender"] = GENDERS.get(record["gender"], "")
    record["appointment_types"] = "; ".join(TYPES[v] for v in data.get("appointment_type", []) if v in TYPES)
    record["name"] = f"{record['first_name']} {record['last_name']}".strip()
    record["booking_id"] = f"APT-{uuid.uuid4().hex[:6].upper()}"
    record["submitted_at"] = datetime.now().isoformat(timespec="seconds")
    return record


def validate(record):
    errors = [f"{LABELS[key]} is required." for key in REQUIRED if not record[key]]
    if record["phone"] and not re.fullmatch(r"\d{10}", record["phone"]):
        errors.append("Phone number must be exactly 10 digits.")
    for key, parse in (("date", date.fromisoformat), ("time", time.fromisoformat)):
        if record[key]:
            try:
                parse(record[key])
            except ValueError:
                errors.append(f"{LABELS[key]} is not valid.")
    return errors


def open_sheet():
    if DATA_FILE.exists():
        workbook = load_workbook(DATA_FILE)
        return workbook, workbook.active

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Appointments"
    sheet.append([title for _, title, _ in COLUMNS])
    for index, (key, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="111111")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"

    contacted = get_column_letter([key for key, _, _ in COLUMNS].index("contacted") + 1)
    choices = DataValidation(type="list", formula1='"Yes,No,No answer"', allow_blank=True)
    choices.add(f"{contacted}2:{contacted}10000")
    sheet.add_data_validation(choices)
    return workbook, sheet


def cell_value(key, value):
    if not value:
        return None
    if key == "submitted_at":
        return datetime.fromisoformat(value)
    if key == "date":
        return date.fromisoformat(value)
    if key == "time":
        return time.fromisoformat(value)
    return value


def write_rows(records):
    workbook, sheet = open_sheet()
    for record in records:
        row = sheet.max_row + 1
        for column, (key, _, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row, column=column, value=cell_value(key, record.get(key, "")))
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
            cell.number_format = FORMATS.get(key, "General")
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("appointment_types", "query", "notes"))
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"

    temp = DATA_FILE.with_name(DATA_FILE.stem + ".tmp.xlsx")
    workbook.save(temp)
    try:
        os.replace(temp, DATA_FILE)
    except PermissionError:
        temp.unlink(missing_ok=True)
        raise


def load_pending():
    if not PENDING_FILE.exists():
        return []
    return json.loads(PENDING_FILE.read_text(encoding="utf-8"))


def save(records):
    with write_lock:
        records = load_pending() + records
        if not records:
            return
        try:
            write_rows(records)
        except PermissionError:
            PENDING_FILE.write_text(json.dumps(records, indent=2), encoding="utf-8")
            print(f"{DATA_FILE.name} is open in another program. {len(records)} booking(s) kept in "
                  f"{PENDING_FILE.name}; they will be added once the file is closed.")
            return
        PENDING_FILE.unlink(missing_ok=True)


def page(title, body):
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{title}</title>
  <link rel="stylesheet" href="/style.css">
</head>
<body>
  <main class="card">
    <header class="card__header">
      <h1>{title}</h1>
    </header>
    <div class="form">
      {body}
    </div>
  </main>
</body>
</html>"""


def confirmation_page(record):
    rows = "".join(
        f"<dt>{LABELS[key]}</dt><dd>{html.escape(record[key]) or '&mdash;'}</dd>"
        for key in LABELS
    )
    name = html.escape(record["first_name"])
    return page("Appointment booked", f"""
      <p class="lead">Thank you, {name}. Your booking ID is <strong>{record['booking_id']}</strong>.
      We will call you on {html.escape(record['phone'])} to confirm.</p>
      <dl class="summary">{rows}</dl>
      <div class="actions"><a class="btn btn--primary" href="/">Book another appointment</a></div>""")


def error_page(errors):
    items = "".join(f"<li>{html.escape(e)}</li>" for e in errors)
    return page("Please check your details", f"""
      <ul class="errors">{items}</ul>
      <div class="actions"><a class="btn btn--primary" href="/">Back to the form</a></div>""")


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def do_GET(self):
        if self._is_public():
            super().do_GET()

    def do_HEAD(self):
        if self._is_public():
            super().do_HEAD()

    def do_POST(self):
        if self.path != "/book":
            self.send_error(HTTPStatus.NOT_FOUND)
            return
        length = int(self.headers.get("Content-Length") or 0)
        if length > MAX_BODY:
            self.send_error(HTTPStatus.REQUEST_ENTITY_TOO_LARGE)
            return
        record = parse_form(self.rfile.read(length).decode("utf-8", errors="replace"))
        errors = validate(record)
        if errors:
            self._respond(HTTPStatus.BAD_REQUEST, error_page(errors))
            return
        save([record])
        self._respond(HTTPStatus.OK, confirmation_page(record))

    def _is_public(self):
        if self.path.split("?", 1)[0] in PUBLIC_FILES:
            return True
        self.send_error(HTTPStatus.NOT_FOUND)
        return False

    def _respond(self, status, body):
        data = body.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


def main():
    save([])
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    url = f"http://localhost:{PORT}/"
    print(f"Appointment form running at {url}")
    print(f"Bookings are saved to {DATA_FILE}")
    print("Press Ctrl+C to stop.")
    webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
